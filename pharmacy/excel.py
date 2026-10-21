"""Strict, previewable XLSX opening inventory import; never imports executable macros."""
from pathlib import Path
from datetime import date, datetime
from hashlib import sha256
from openpyxl import Workbook, load_workbook
from .core import money, integer, valid_date, now
HEADERS=['sku','name','barcode','category','selling_price','reorder_level','batch','expiry','opening_quantity','unit_cost']
def template(path):
    w=Workbook(); ws=w.active; ws.title='Inventory'; ws.append(HEADERS)
    ws.freeze_panes='A2'; ws.auto_filter.ref='A1:J1'
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=22
    w.save(path)
def preview(path):
    path=Path(path)
    if path.suffix.lower()!='.xlsx': raise ValueError('Use an .xlsx workbook; save legacy .xls files as .xlsx first')
    if path.stat().st_size>20*1024*1024: raise ValueError('Workbook exceeds 20 MB')
    data=path.read_bytes(); digest=sha256(data).hexdigest()
    from io import BytesIO
    w=load_workbook(BytesIO(data),read_only=True,data_only=False)
    try:
        ws=w.active
        if ws.max_row>50001: raise ValueError('Maximum 50,000 rows per import')
        rows=iter(ws.iter_rows(values_only=True))
        if list(next(rows,()))!=HEADERS: raise ValueError('Columns must exactly match the generated template')
        clean=[]; errors=[]
        for number,row in enumerate(rows,2):
            if all(v is None for v in row): continue
            try:
                if any(isinstance(v,str) and v.startswith('=') for v in row): raise ValueError('Formulas are not accepted; paste values')
                r=dict(zip(HEADERS,row))
                for key in ('sku','name','batch'):
                    if not str(r[key] or '').strip(): raise ValueError(key+' required')
                for key in ('sku','name','barcode','category','batch'): r[key]=str(r[key] or '').strip()
                r['price']=money(r.pop('selling_price')); r['cost']=money(r.pop('unit_cost'))
                r['quantity']=integer(r.pop('opening_quantity'),1); r['reorder']=integer(r.pop('reorder_level'))
                expiry=r['expiry']
                if isinstance(expiry,(datetime,date)): expiry=expiry.date().isoformat() if isinstance(expiry,datetime) else expiry.isoformat()
                r['expiry']=valid_date(expiry)
                if r['expiry']<date.today().isoformat(): raise ValueError('Opening stock must be unexpired')
                clean.append(r)
            except (ValueError,TypeError) as e: errors.append(f'Row {number}: {e}')
        if errors: raise ValueError('\n'.join(errors[:30]))
        if not clean: raise ValueError('No inventory rows found')
        # Multiple batches allowed only with consistent product details.
        products={}; barcodes={}; batch_keys=set()
        for r in clean:
            key=(r['sku'],r['batch'],r['expiry'])
            if key in batch_keys: raise ValueError('Duplicate SKU/batch/expiry in workbook: '+r['sku'])
            batch_keys.add(key)
            identity=tuple(r[k] for k in ('name','barcode','category','price','reorder'))
            if r['sku'] in products and products[r['sku']]!=identity: raise ValueError('Conflicting product details: '+r['sku'])
            products[r['sku']]=identity
            if r['barcode']:
                if r['barcode'] in barcodes and barcodes[r['barcode']]!=r['sku']: raise ValueError('Barcode used by multiple SKUs')
                barcodes[r['barcode']]=r['sku']
        return clean,digest
    finally: w.close()
def apply(store,clean,digest):
    store.authorize(True)
    with store.transaction():
        if store.rows('SELECT id FROM imports WHERE digest=?',(digest,)): raise ValueError('This exact workbook was already imported')
        ids={}
        for r in clean:
            if r['sku'] not in ids:
                if store.rows('SELECT id FROM products WHERE sku=?',(r['sku'],)): raise ValueError('SKU already exists; opening import only creates new products: '+r['sku'])
                ids[r['sku']]=store.db.execute('INSERT INTO products(sku,name,barcode,category,price,reorder) VALUES(?,?,?,?,?,?)',(r['sku'],r['name'],r['barcode'] or None,r['category'],r['price'],r['reorder'])).lastrowid
            store.db.execute('INSERT INTO batches(product_id,batch,expiry,quantity,cost,original_quantity) VALUES(?,?,?,?,?,?)',(ids[r['sku']],r['batch'],r['expiry'],r['quantity'],r['cost'],r['quantity']))
        store.db.execute('INSERT INTO imports(digest,created,rows) VALUES(?,?,?)',(digest,now(),len(clean)))
        store.audit('Opening inventory import',f'{len(clean)} batches; SHA256 {digest}')
